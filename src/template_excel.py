#!/usr/bin/env python3
"""Buat templat Excel pengisian data Dalops untuk Anper berbasis gas
(PKT, PIM, PSP, PKC), mengikuti kontrak data dashboard CAK PRI v3 (butir 11.2).

- Nama sheet dan nama kolom baris 1 = kontrak data; berkas yang sudah diisi
  bisa langsung dimuat lewat tombol "Muat Data" tanpa pemetaan tambahan.
- Kolom yang bisa dihitung (jenis, stok awal, stok akhir, total gas,
  waktu layar, total RKAP) berisi rumus; yang diisi Anper hanya sel kuning.
- Kunci baris (tanggal, material, unit, pemasok) sudah disiapkan untuk
  1 Jan – 31 Des 2026. Nama unit/pemasok diambil dari data DUMMY sebagai
  contoh dan HARUS diganti dengan nama sebenarnya (lihat sheet PANDUAN).

Pakai:  python3 src/template_excel.py   (dari akar repo) → templates/*.xlsx
"""
import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AKAR = Path(__file__).resolve().parent.parent
DUMMY = json.loads((AKAR / 'data' / 'DUMMY_CAK_PRI_PIGroup_data.json').read_text(encoding='utf-8'))['entitas']
KELUAR = AKAR / 'templates'

TAHUN = 2026
HARI = [dt.date(TAHUN, 1, 1) + dt.timedelta(days=i) for i in range(365)]
POSISI = "profil!$B$9"                   # sel posisi_data di sheet profil

F = 'Arial'
HEAD = dict(font=Font(name=F, bold=True, color='FFFFFF', size=10), fill=PatternFill('solid', fgColor='1F5F3A'),
            alignment=Alignment(horizontal='center', vertical='center', wrap_text=True))
ISI = dict(font=Font(name=F, color='0000FF', size=10), fill=PatternFill('solid', fgColor='FFF2CC'))
RUMUS = dict(font=Font(name=F, color='000000', size=10), fill=PatternFill('solid', fgColor='EDEDED'))
KUNCI = dict(font=Font(name=F, color='000000', size=10))
TIPIS = Side(style='thin', color='D0D7D2')
FMT_TGL, FMT_T, FMT_G, FMT_R = 'dd-mmm-yyyy', '#,##0.0;-#,##0.0;0', '#,##0.00', '0.000'


def gaya(c, g, fmt=None):
    for k, v in g.items():
        setattr(c, k, v)
    if fmt:
        c.number_format = fmt


def lembar(wb, nama, kolom, lebar, ket):
    """Buat sheet kontrak: header baris 1 + komentar penjelasan per kolom."""
    ws = wb.create_sheet(nama)
    for j, (k, w) in enumerate(zip(kolom, lebar), 1):
        c = ws.cell(1, j, k)
        gaya(c, HEAD)
        if ket.get(k):
            c.comment = Comment(ket[k], 'CAK PRI')
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30
    ws.sheet_view.zoomScale = 90
    return ws


def baris(ws, r, nilai):
    """nilai: list of (isi, gaya, format)."""
    for j, (v, g, fmt) in enumerate(nilai, 1):
        c = ws.cell(r, j, v)
        gaya(c, g, fmt)


def jenis(r):
    return (f'=IF(A{r}="","",IF(A{r}<={POSISI},"Realisasi","Rencana"))', RUMUS, None)


def validasi(ws, rentang, pilihan, judul):
    dv = DataValidation(type='list', formula1='"' + ','.join(pilihan) + '"', allow_blank=True,
                        errorTitle='Isian tidak dikenal', error=f'Pilih salah satu: {", ".join(pilihan)}',
                        promptTitle=judul, prompt='Pilih dari daftar', showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(rentang)


def rows(e, s):
    t = DUMMY[e][s]
    return [dict(zip(t['kolom'], r)) for r in t['baris']]


def buat(kode):
    prof = {r['parameter']: r['nilai'] for r in rows(kode, 'profil')}
    pemasok = list(dict.fromkeys(r['pemasok'] for r in rows(kode, 'harian_gas')))
    unit = list(dict.fromkeys((r['unit'], r['produk']) for r in rows(kode, 'harian_produksi')))
    produk_l1 = list(dict.fromkeys(r['produk'] for r in rows(kode, 'stok_produk_lini1')))
    sumber = rows(kode, 'sumber_pasokan')
    MATERIAL = [('NH3', 'Amoniak', 'Produksi sendiri (gas alam)', 'ton', 'Pemenuhan Safety Stock Amoniak (penyangga gas)'),
                ('KCL', 'KCl', 'Impor', 'ton', 'Pemenuhan Safety Stock Bahan Baku Non-Gas'),
                ('DAP', 'DAP', 'Impor', 'ton', 'Pemenuhan Safety Stock Bahan Baku Non-Gas'),
                ('ZA', 'ZA', 'Domestik (antar-entitas PI Group)', 'ton', 'Pemenuhan Safety Stock Bahan Baku Non-Gas'),
                ('CLAY', 'Clay / filler', 'Domestik', 'ton', 'Pemenuhan Safety Stock Bahan Baku Non-Gas')]

    wb = Workbook()
    wb.remove(wb.active)
    pan = wb.create_sheet('PANDUAN')                     # diisi paling akhir

    # ---------------- profil ----------------
    ws = lembar(wb, 'profil', ['parameter', 'nilai', 'keterangan'], [24, 46, 70], {})
    P = [('kode', kode, KUNCI, None, 'Jangan diubah. Dipakai dashboard untuk mengenali entitas.'),
         ('nama', prof['nama'], ISI, None, 'Nama resmi perusahaan.'),
         ('lokasi', prof['lokasi'], ISI, None, 'Kota, provinsi lokasi pabrik.'),
         ('pelabuhan_bongkar', prof['pelabuhan_bongkar'], ISI, None, 'Pelabuhan bongkar impor. Periksa: nama dari data dummy.'),
         ('pelabuhan_lat', prof['pelabuhan_lat'], ISI, '0.00', 'Lintang pelabuhan bongkar (derajat desimal, selatan negatif).'),
         ('pelabuhan_lon', prof['pelabuhan_lon'], ISI, '0.00', 'Bujur pelabuhan bongkar (derajat desimal).'),
         ('warna_aksen', None, ISI, None, 'Kode warna brand, mis. #F5A623. Boleh dikosongkan.'),
         ('posisi_data', None, ISI, FMT_TGL, 'WAJIB. Tanggal realisasi terakhir. Kolom "jenis" di semua sheet mengikuti tanggal ini.'),
         ('tanggal_update', None, ISI, FMT_TGL, 'Tanggal berkas ini disusun.'),
         ('periode_konsumsi',
          '=IF(B9="","",DAY(B9-29)&" "&CHOOSE(MONTH(B9-29),"Jan","Feb","Mar","Apr","Mei","Jun","Jul","Agu","Sep","Okt","Nov","Des")'
          '&" – "&DAY(B9)&" "&CHOOSE(MONTH(B9),"Jan","Feb","Mar","Apr","Mei","Jun","Jul","Agu","Sep","Okt","Nov","Des")&" "&YEAR(B9)&" (rata-rata 30 hari)")',
          RUMUS, None, 'Otomatis dari posisi_data.'),
         ('konteks', None, ISI, None, 'Satu kalimat konteks risiko utama saat ini (mis. keandalan pasokan gas hulu, jadwal turnaround).'),
         ('dosis_urea_kg_per_ha', None, ISI, '#,##0', 'Dosis urea per hektare sesuai rekomendasi yang berlaku.'),
         ('dosis_npk_kg_per_ha', None, ISI, '#,##0', 'Dosis NPK per hektare sesuai rekomendasi yang berlaku.'),
         ('produk_npk', None, ISI, None, 'Nama & formula NPK utama, mis. NPK 15-15-15.'),
         ('status_data', 'ASLI', KUNCI, None, 'Biarkan "ASLI". Kata DUMMY di sel ini membuat dashboard memasang pita DATA DUMMY.')]
    for i, (k, v, g, fmt, ket) in enumerate(P, 2):
        baris(ws, i, [(k, KUNCI, None), (v, g, fmt), (ket, dict(font=Font(name=F, size=9, color='595959')), None)])
        ws.cell(i, 3).alignment = Alignment(wrap_text=True, vertical='top')

    # ---------------- harian_gas ----------------
    ws = lembar(wb, 'harian_gas', ['tanggal', 'jenis', 'pemasok', 'kontrak_bbtud', 'pasokan_bbtud'], [14, 12, 36, 15, 15],
                {'jenis': 'Otomatis: Realisasi bila tanggal ≤ posisi_data, selain itu Rencana.',
                 'pemasok': 'Nama pemasok/lapangan gas. Harus sama persis dengan sheet sumber_pasokan.',
                 'kontrak_bbtud': 'Otomatis dari volume kontrak di sheet sumber_pasokan. Boleh ditimpa bila kontrak berubah di tengah tahun.',
                 'pasokan_bbtud': 'Realisasi: pasokan aktual harian. Rencana: nominasi/rencana pasokan.'})
    r = 2
    for d in HARI:
        for p in pemasok:
            baris(ws, r, [(d, KUNCI, FMT_TGL), jenis(r), (p, KUNCI, None),
                          (f'=IFERROR(INDEX(sumber_pasokan!$E:$E,MATCH(C{r},sumber_pasokan!$B:$B,0)),"")', RUMUS, FMT_G),
                          (None, ISI, FMT_G)])
            r += 1
    n_gas = r - 1

    # ---------------- ringkasan_gas ----------------
    ws = lembar(wb, 'ringkasan_gas', ['tanggal', 'jenis', 'kontrak_total_bbtud', 'pasokan_total_bbtud', 'kebutuhan_rencana_bbtud',
                                      'pemakaian_aktual_bbtud', 'gas_per_ton_amoniak_mmbtu', 'gas_per_ton_urea_mmbtu'],
                [14, 12, 15, 15, 17, 17, 17, 17],
                {'kontrak_total_bbtud': 'Otomatis: jumlah kontrak semua pemasok di harian_gas.',
                 'pasokan_total_bbtud': 'Otomatis: jumlah pasokan semua pemasok di harian_gas.',
                 'kebutuhan_rencana_bbtud': 'Kebutuhan gas operasi menurut rencana produksi (dasar KRI Pemenuhan Pasokan Gas).',
                 'pemakaian_aktual_bbtud': 'Pemakaian gas aktual. Kosongkan pada baris Rencana.',
                 'gas_per_ton_amoniak_mmbtu': 'Konsumsi spesifik gas per ton amoniak (bahan baku + bahan bakar).',
                 'gas_per_ton_urea_mmbtu': 'Konsumsi spesifik gas per ton urea (utilitas).'})
    for r, d in enumerate(HARI, 2):
        baris(ws, r, [(d, KUNCI, FMT_TGL), jenis(r),
                      (f'=SUMIFS(harian_gas!$D:$D,harian_gas!$A:$A,A{r})', RUMUS, FMT_G),
                      (f'=SUMIFS(harian_gas!$E:$E,harian_gas!$A:$A,A{r})', RUMUS, FMT_G),
                      (None, ISI, FMT_G), (None, ISI, FMT_G), (None, ISI, FMT_G), (None, ISI, FMT_G)])

    # ---------------- harian_bahan_baku ----------------
    ws = lembar(wb, 'harian_bahan_baku', ['tanggal', 'jenis', 'kode_material', 'material', 'stok_awal_t', 'kedatangan_t',
                                          'produksi_sendiri_t', 'pemakaian_t', 'dijual_transfer_t', 'stok_akhir_t', 'cek_selisih_t'],
                [14, 12, 13, 15, 14, 14, 15, 14, 15, 14, 13],
                {'stok_awal_t': 'Isi hanya pada 1 Januari. Hari berikutnya otomatis = stok akhir hari sebelumnya.',
                 'kedatangan_t': 'Bongkaran kapal/kiriman yang masuk gudang hari itu (ton).',
                 'produksi_sendiri_t': 'Hanya amoniak: produksi pabrik amoniak hari itu (ton).',
                 'pemakaian_t': 'Pemakaian oleh pabrik (amoniak ke urea; bahan NPK ke pabrik NPK).',
                 'dijual_transfer_t': 'Penjualan atau transfer ke entitas lain (ton).',
                 'stok_akhir_t': 'Otomatis = awal + kedatangan + produksi − pemakaian − jual/transfer. Pada baris Realisasi boleh ditimpa dengan stok hasil opname.',
                 'cek_selisih_t': 'Kontrol: selisih neraca. Harus 0 kecuali stok akhir ditimpa angka opname. Tidak dibaca dashboard.'})
    r = 2
    for kode_m, nama_m, *_ in MATERIAL:
        for i, d in enumerate(HARI):
            awal = (None, ISI, FMT_T) if i == 0 else (f'=J{r-1}', RUMUS, FMT_T)
            baris(ws, r, [(d, KUNCI, FMT_TGL), jenis(r), (kode_m, KUNCI, None), (nama_m, KUNCI, None), awal,
                          (None, ISI, FMT_T), ((None, ISI, FMT_T) if kode_m == 'NH3' else (0, KUNCI, FMT_T)),
                          (None, ISI, FMT_T), (None, ISI, FMT_T),
                          (f'=E{r}+F{r}+G{r}-H{r}-I{r}', RUMUS, FMT_T),
                          (f'=ROUND(E{r}+F{r}+G{r}-H{r}-I{r}-J{r},1)', RUMUS, FMT_T)])
            r += 1
    n_bb = r - 1

    # ---------------- harian_produksi ----------------
    ws = lembar(wb, 'harian_produksi', ['tanggal', 'jenis', 'unit', 'produk', 'kapasitas_t_per_hari', 'produksi_t'], [14, 12, 16, 12, 16, 14],
                {'unit': 'Nama unit pabrik. Contoh dari data dummy — ganti dengan nama unit sebenarnya (sama di semua sheet).',
                 'produk': 'Amoniak, Urea, atau NPK (huruf persis seperti ini).',
                 'kapasitas_t_per_hari': 'Isi pada baris pertama tiap unit; baris berikutnya mengikuti otomatis (boleh ditimpa).',
                 'produksi_t': 'Realisasi produksi harian, atau rencana produksi untuk baris Rencana (ton).'})
    r = 2
    for u, p in unit:
        for i, d in enumerate(HARI):
            kap = (None, ISI, '#,##0') if i == 0 else (f'=E{r-1}', RUMUS, '#,##0')
            baris(ws, r, [(d, KUNCI, FMT_TGL), jenis(r), (u, KUNCI, None), (p, KUNCI, None), kap, (None, ISI, FMT_T)])
            r += 1
    validasi(ws, f'D2:D{r-1}', ['Amoniak', 'Urea', 'NPK'], 'Produk')

    # ---------------- rkap_bulanan ----------------
    bln = ['jan', 'feb', 'mar', 'apr', 'mei', 'jun', 'jul', 'agu', 'sep', 'okt', 'nov', 'des']
    ws = lembar(wb, 'rkap_bulanan', ['unit', 'produk'] + bln + ['total_2026'], [16, 11] + [11] * 12 + [13],
                {'total_2026': 'Otomatis: jumlah Januari–Desember.', 'unit': 'Sama persis dengan nama unit di harian_produksi.'})
    for r, (u, p) in enumerate(unit, 2):
        baris(ws, r, [(u, KUNCI, None), (p, KUNCI, None)] + [(None, ISI, '#,##0')] * 12 + [(f'=SUM(C{r}:N{r})', RUMUS, '#,##0')])

    # ---------------- stok_produk_lini1 ----------------
    ws = lembar(wb, 'stok_produk_lini1', ['tanggal', 'jenis', 'produk', 'stok_awal_t', 'masuk_t', 'penyaluran_t', 'stok_akhir_t',
                                          'curah_di_pabrik_t', 'cek_selisih_t'], [14, 12, 22, 14, 13, 14, 14, 16, 13],
                {'stok_awal_t': 'Isi hanya pada 1 Januari; hari berikutnya otomatis.',
                 'masuk_t': 'Produk masuk gudang Lini I dari pabrik (ton).',
                 'penyaluran_t': 'Produk keluar/disalurkan dari Lini I (ton).',
                 'stok_akhir_t': 'Otomatis; pada baris Realisasi boleh ditimpa dengan stok fisik.',
                 'curah_di_pabrik_t': 'Produk curah di pabrik di luar gudang kantong (ton).',
                 'cek_selisih_t': 'Kontrol neraca, tidak dibaca dashboard.'})
    r = 2
    for p in produk_l1:
        for i, d in enumerate(HARI):
            awal = (None, ISI, FMT_T) if i == 0 else (f'=G{r-1}', RUMUS, FMT_T)
            baris(ws, r, [(d, KUNCI, FMT_TGL), jenis(r), (p, KUNCI, None), awal, (None, ISI, FMT_T), (None, ISI, FMT_T),
                          (f'=D{r}+E{r}-F{r}', RUMUS, FMT_T), (None, ISI, FMT_T), (f'=ROUND(D{r}+E{r}-F{r}-G{r},1)', RUMUS, FMT_T)])
            r += 1

    # ---------------- jadwal_kapal ----------------
    ws = lembar(wb, 'jadwal_kapal', ['material', 'negara_asal', 'pelabuhan_muat', 'nama_kapal', 'volume_t', 'eta_terkini', 'eta_awal',
                                     'status', 'jarak_nm', 'waktu_layar_hari', 'lewat_hormuz', 'catatan'],
                [11, 18, 26, 24, 12, 14, 14, 16, 11, 14, 13, 40],
                {'material': 'KCl, DAP, ZA, atau Amoniak.', 'eta_terkini': 'ETA menurut rencana terbaru.',
                 'eta_awal': 'ETA pada rencana awal. Selisih dengan ETA terkini = keterlambatan.',
                 'jarak_nm': 'Jarak laut pelabuhan muat → pelabuhan bongkar (nautical mile).',
                 'waktu_layar_hari': 'Otomatis: jarak ÷ (12,5 knot × 24 jam).',
                 'lewat_hormuz': 'Ya bila rute melewati Selat Hormuz (mis. dari Teluk Persia).'})
    for r in range(2, 62):
        baris(ws, r, [(None, ISI, None), (None, ISI, None), (None, ISI, None), (None, ISI, None), (None, ISI, '#,##0'),
                      (None, ISI, FMT_TGL), (None, ISI, FMT_TGL), (None, ISI, None), (None, ISI, '#,##0'),
                      (f'=IF(I{r}="","",ROUND(I{r}/(12.5*24),1))', RUMUS, '0.0'), (None, ISI, None), (None, ISI, None)])
    validasi(ws, 'A2:A61', ['KCl', 'DAP', 'ZA', 'Amoniak'], 'Material')
    validasi(ws, 'H2:H61', ['Tiba', 'Dalam pelayaran', 'Rencana', 'Terlambat'], 'Status kapal')
    validasi(ws, 'K2:K61', ['Ya', 'Tidak'], 'Lewat Hormuz')

    # ---------------- perubahan_versi ----------------
    ws = lembar(wb, 'perubahan_versi', ['versi_lama', 'versi_baru', 'material', 'nama_kapal', 'eta_lama', 'eta_baru', 'volume_lama_t',
                                        'volume_baru_t', 'alasan'], [13, 13, 11, 24, 13, 13, 14, 14, 44],
                {'versi_lama': 'Tanggal versi master sebelumnya.', 'versi_baru': 'Tanggal versi master ini.',
                 'alasan': 'Sebab perubahan ETA atau volume.'})
    for r in range(2, 32):
        baris(ws, r, [(None, ISI, FMT_TGL), (None, ISI, FMT_TGL), (None, ISI, None), (None, ISI, None), (None, ISI, FMT_TGL),
                      (None, ISI, FMT_TGL), (None, ISI, '#,##0'), (None, ISI, '#,##0'), (None, ISI, None)])
    validasi(ws, 'C2:C31', ['KCl', 'DAP', 'ZA', 'Amoniak'], 'Material')

    # ---------------- status_unit ----------------
    ws = lembar(wb, 'status_unit', ['unit', 'status', 'mulai', 'selesai', 'keterangan'], [16, 14, 14, 14, 46],
                {'status': 'Beroperasi, Berhenti, atau Beban turun — pada tanggal posisi data.',
                 'mulai': 'Untuk unit berhenti: tanggal mulai.', 'selesai': 'Untuk unit berhenti: perkiraan tanggal kembali beroperasi.',
                 'keterangan': 'Mis. turnaround terjadwal, trip kompresor, menunggu suku cadang.'})
    for r, (u, _) in enumerate(unit, 2):
        baris(ws, r, [(u, KUNCI, None), ('Beroperasi', ISI, None), (None, ISI, FMT_TGL), (None, ISI, FMT_TGL), (None, ISI, None)])
    validasi(ws, f'B2:B{len(unit)+1}', ['Beroperasi', 'Berhenti', 'Beban turun'], 'Status unit')

    # ---------------- parameter_kri ----------------
    ws = lembar(wb, 'parameter_kri', ['kode_material', 'material', 'sifat_pasokan', 'satuan', 'safety_stock', 'deadstock',
                                      'ambang_aman_pct', 'ambang_bahaya_pct', 'nama_kri'], [13, 15, 30, 9, 13, 12, 15, 16, 62],
                {'safety_stock': 'Sesuai Kesepakatan Safety Stock Bahan Baku PI Group. Dashboard tetap memakai angka Kesepakatan di profil entitas; bila berbeda, dashboard memberi peringatan.',
                 'deadstock': 'Sesuai data deadstock PI Group.',
                 'ambang_aman_pct': 'KRI non-gas: Aman ≥ 100%.', 'ambang_bahaya_pct': 'KRI non-gas: Bahaya < 92,25%.'})
    for r, (km, nm, sf, sat, kri) in enumerate(MATERIAL, 2):
        baris(ws, r, [(km, KUNCI, None), (nm, KUNCI, None), (sf, KUNCI, None), (sat, KUNCI, None), (None, ISI, '#,##0'),
                      (None, ISI, '#,##0'), (100, KUNCI, '0.00'), (92.25, KUNCI, '0.00'), (kri, KUNCI, None)])
    r = len(MATERIAL) + 2
    baris(ws, r, [('GAS', KUNCI, None), ('Gas alam', KUNCI, None), ('Pipa (kontrak PJBG)', KUNCI, None), ('BBTUD', KUNCI, None),
                  (None, KUNCI, None), (None, KUNCI, None), (95, ISI, '0.00'), (90, ISI, '0.00'),
                  ('Pemenuhan Pasokan Gas terhadap Kebutuhan Operasi (ambang usulan, belum resmi)', KUNCI, None)])
    ws.cell(r, 7).comment = Comment('Ambang usulan. Ganti bila ambang resmi KRI gas sudah ditetapkan.', 'CAK PRI')

    # ---------------- cons_rate ----------------
    ws = lembar(wb, 'cons_rate', ['produk', 'bahan', 'rasio', 'satuan'], [18, 38, 11, 13],
                {'produk': 'Untuk NPK, tulis diawali "NPK", mis. NPK 15-15-15.', 'rasio': 'Rasio konsumsi menurut Dalops.'})
    CR = [('Urea', 'Amoniak', 'ton/ton'), ('Urea', 'Gas alam (utilitas)', 'MMBTU/ton'),
          ('Amoniak', 'Gas alam (bahan baku + bahan bakar)', 'MMBTU/ton'),
          ('NPK 15-15-15', 'DAP', 'ton/ton'), ('NPK 15-15-15', 'KCl', 'ton/ton'), ('NPK 15-15-15', 'Urea (internal untuk NPK)', 'ton/ton'),
          ('NPK 15-15-15', 'ZA', 'ton/ton'), ('NPK 15-15-15', 'Clay / filler', 'ton/ton')]
    for r, (p, b, s) in enumerate(CR, 2):
        baris(ws, r, [(p, ISI if p.startswith('NPK') else KUNCI, None), (b, KUNCI, None), (None, ISI, FMT_R), (s, KUNCI, None)])

    # ---------------- sumber_pasokan ----------------
    ws = lembar(wb, 'sumber_pasokan', ['id', 'nama', 'material', 'moda', 'volume_acuan', 'satuan_volume'], [11, 36, 15, 9, 14, 20],
                {'nama': 'Untuk gas: nama pemasok/lapangan persis seperti di harian_gas. Untuk impor: "Impor <bahan> <negara>" agar Simulator bisa mencocokkan kargo.',
                 'volume_acuan': 'Gas: kontrak BBTUD. Impor: ton per kargo.'})
    for r, s in enumerate(sumber, 2):
        baris(ws, r, [(s['id'], KUNCI, None), (s['nama'], ISI, None), (s['material'], KUNCI, None), (s['moda'], KUNCI, None),
                      (None, ISI, '#,##0.00'), (s['satuan_volume'], KUNCI, None)])
    for r in range(len(sumber) + 2, len(sumber) + 8):
        baris(ws, r, [(None, ISI, None)] * 4 + [(None, ISI, '#,##0.00'), (None, ISI, None)])

    # ---------------- transfer_antar_entitas ----------------
    ws = lembar(wb, 'transfer_antar_entitas', ['tanggal', 'jenis', 'dari', 'ke', 'material', 'volume_t'], [14, 12, 9, 9, 12, 13],
                {'jenis': 'Realisasi atau Rencana.', 'dari': 'Kode entitas pengirim.', 'ke': 'Kode entitas penerima.'})
    for r in range(2, 42):
        baris(ws, r, [(None, ISI, FMT_TGL), (None, ISI, None), (None, ISI, None), (None, ISI, None), (None, ISI, None), (None, ISI, '#,##0')])
    validasi(ws, 'B2:B41', ['Realisasi', 'Rencana'], 'Jenis')
    validasi(ws, 'C2:D41', ['PKG', 'PKT', 'PIM', 'PSP', 'PKC'], 'Entitas')
    validasi(ws, 'E2:E41', ['Amoniak', 'Urea', 'ZA', 'KCl', 'DAP'], 'Material')

    # ---------------- ambang_siaga ----------------
    ws = lembar(wb, 'ambang_siaga', ['tingkat', 'pemicu', 'tindakan'], [20, 70, 80],
                {'pemicu': 'USULAN dari tahap pengembangan. Sesuaikan dengan ambang yang disetujui.'})
    for r, s in enumerate(rows(kode, 'ambang_siaga'), 2):
        baris(ws, r, [(s['tingkat'], KUNCI, None), (s['pemicu'], ISI, None), (s['tindakan'], ISI, None)])
        for j in (2, 3):
            ws.cell(r, j).alignment = Alignment(wrap_text=True, vertical='top')

    # ---------------- PANDUAN ----------------
    ws = pan
    ws.sheet_view.showGridLines = False
    for col, w in zip('ABCDEFGHIJK', [3, 28, 62, 22, 22, 14, 14, 14, 14, 14, 14]):
        ws.column_dimensions[col].width = w
    t = lambda r, c, v, **k: gaya(ws.cell(r, c, v), {'font': Font(name=F, **k)})
    t(1, 2, f'Templat Pengisian Data Dalops 2026 · {kode}', bold=True, size=16, color='1F5F3A')
    t(2, 2, f'{prof["nama"]} · untuk dashboard CAK PRI PI Group v3', size=11, color='595959')
    t(4, 2, 'Arti warna sel', bold=True, size=11)
    for i, (g, teks) in enumerate([(ISI, 'Kuning, teks biru: DIISI oleh Anper'),
                                   (RUMUS, 'Abu-abu: rumus otomatis — jangan diubah (kecuali disebut "boleh ditimpa")'),
                                   (KUNCI, 'Putih: kunci baris (tanggal, material, unit) — sudah disiapkan')]):
        c = ws.cell(5 + i, 2, 'contoh')
        gaya(c, g)
        c.border = Border(TIPIS, TIPIS, TIPIS, TIPIS)
        t(5 + i, 3, teks, size=10)

    t(9, 2, 'Langkah pengisian', bold=True, size=11)
    langkah = [
        '1. Sheet profil: isi posisi_data (tanggal realisasi terakhir), tanggal_update, dosis, produk NPK, konteks. Kolom "jenis" di semua sheet otomatis mengikuti posisi_data.',
        '2. Ganti nama unit, pemasok gas, dan sumber pasokan contoh dengan nama sebenarnya (Ctrl+H, ganti di seluruh workbook agar konsisten antarsheet).',
        '3. Isi sheet harian: angka Realisasi sampai posisi_data, lalu angka Rencana sampai 31 Desember. Stok awal cukup diisi pada 1 Januari.',
        '4. Isi jadwal_kapal untuk semua kargo 2026 (tiba, dalam pelayaran, rencana). Isi perubahan_versi bila ETA/volume berubah dari versi sebelumnya.',
        '5. Isi status_unit (posisi hari ini), parameter_kri (safety stock & deadstock Kesepakatan PI Group), cons_rate, sumber_pasokan.',
        '6. Periksa tabel "Cek otomatis" di bawah: semua harus 0 / OK sebelum dikirim.',
        f'7. Simpan sebagai .xlsx dengan nama mengandung kode entitas, mis. Dalops_{kode}_2026_posisi_<tanggal>.xlsx, lalu muat lewat tombol "Muat Data" di dashboard.',
        'Jangan mengganti nama sheet, nama kolom di baris 1, atau menyisipkan baris di atas header. Sheet PANDUAN dan kolom cek_selisih_t tidak dibaca dashboard.']
    for i, s in enumerate(langkah):
        c = ws.cell(10 + i, 2, s)
        gaya(c, {'font': Font(name=F, size=10)})
        ws.merge_cells(start_row=10 + i, start_column=2, end_row=10 + i, end_column=8)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[10 + i].height = 28

    r0 = 19
    t(r0, 2, 'Isi tiap sheet', bold=True, size=11)
    SH = [('profil', 'Identitas, posisi data, dosis, konteks', 'Semua modul'),
          ('harian_gas', 'Pasokan harian per pemasok (BBTUD) vs kontrak', '01, 08, 09'),
          ('ringkasan_gas', 'Total pasokan (otomatis), kebutuhan, pemakaian, konsumsi spesifik', '01, 02, 05, 09'),
          ('harian_bahan_baku', 'Neraca harian amoniak, KCl, DAP, ZA, clay', '01, 05, 06, 08, 09'),
          ('harian_produksi', 'Produksi harian per unit amoniak/urea/NPK', '02, 08, 09, 11, 12'),
          ('rkap_bulanan', 'RKAP 2026 per unit per bulan', '02'),
          ('stok_produk_lini1', 'Stok, masuk, penyaluran produk di gudang Lini I', '03'),
          ('jadwal_kapal', 'Seluruh kargo 2026 beserta ETA awal & terkini', '01, 04, 05, 06'),
          ('perubahan_versi', 'Perubahan ETA/volume antar-versi master', '07'),
          ('status_unit', 'Unit berhenti/turnaround/trip pada posisi data', '01, 02, 13'),
          ('parameter_kri', 'Safety stock & deadstock (Kesepakatan PI Group), ambang KRI', '01, 13'),
          ('cons_rate', 'Rasio konsumsi bahan per ton produk', '02, 09, 10, 11, 12'),
          ('sumber_pasokan', 'Daftar sumber untuk Simulator; kontrak gas', '08'),
          ('transfer_antar_entitas', 'Kiriman antar-entitas PI Group', '11, 00'),
          ('ambang_siaga', 'Pemicu Siaga 1–3 dan tindakannya', '13')]
    for j, h in enumerate(['Sheet', 'Isi', 'Modul dashboard'], 2):
        gaya(ws.cell(r0 + 1, j, h), HEAD)
    for i, (a, b, c_) in enumerate(SH):
        for j, v in enumerate([a, b, c_], 2):
            c = ws.cell(r0 + 2 + i, j, v)
            gaya(c, {'font': Font(name=F, size=10, bold=(j == 2))})
            c.border = Border(bottom=TIPIS)

    r1 = r0 + 3 + len(SH)
    t(r1, 2, 'Contoh pengisian (format yang diharapkan; angka contoh)', bold=True, size=11)
    CONTOH = [('harian_bahan_baku', ['tanggal', 'jenis', 'kode_material', 'material', 'stok_awal_t', 'kedatangan_t', 'produksi_sendiri_t', 'pemakaian_t', 'dijual_transfer_t', 'stok_akhir_t'],
               [dt.date(TAHUN, 10, 8), 'Realisasi', 'KCL', 'KCl', 14467.6, 0, 0, 217.6, 0, 14250.0]),
              ('harian_gas', ['tanggal', 'jenis', 'pemasok', 'kontrak_bbtud', 'pasokan_bbtud'],
               [dt.date(TAHUN, 10, 8), 'Realisasi', pemasok[0], 150, 115.7]),
              ('jadwal_kapal', ['material', 'negara_asal', 'pelabuhan_muat', 'nama_kapal', 'volume_t', 'eta_terkini', 'eta_awal', 'status', 'jarak_nm', 'waktu_layar_hari', 'lewat_hormuz'],
               ['KCl', 'Kanada', 'Vancouver', 'MV Nama Kapal', 31400, dt.date(TAHUN, 10, 19), dt.date(TAHUN, 10, 3), 'Terlambat', 8150, 27.2, 'Tidak'])]
    rr = r1 + 1
    for nama, hdr, nilai in CONTOH:
        t(rr, 2, nama, bold=True, size=10, color='1F5F3A')
        rr += 1
        for j, (h, v) in enumerate(zip(hdr, nilai), 2):
            gaya(ws.cell(rr, j, h), {'font': Font(name=F, size=9, bold=True, color='595959')})
            c = ws.cell(rr + 1, j, v)
            gaya(c, {'font': Font(name=F, size=10, italic=True)}, FMT_TGL if isinstance(v, dt.date) else None)
        rr += 3
    t(rr - 1, 2, 'Baris contoh di atas hanya ilustrasi format, tidak dibaca dashboard.', size=9, italic=True, color='7F7F7F')

    r2 = rr + 1
    t(r2, 2, 'Cek otomatis', bold=True, size=11)
    CEK = [('posisi_data terisi', '=IF(profil!B9="","BELUM","OK")'),
           ('Baris Realisasi bahan baku tanpa angka pemakaian', '=COUNTIFS(harian_bahan_baku!B:B,"Realisasi",harian_bahan_baku!H:H,"")'),
           ('Baris neraca bahan baku dengan selisih > 1 t', f'=SUMPRODUCT(--(ABS(harian_bahan_baku!K2:K{n_bb})>1))'),
           ('Baris Realisasi gas tanpa angka pasokan', '=COUNTIFS(harian_gas!B:B,"Realisasi",harian_gas!E:E,"")'),
           ('Hari tanpa kebutuhan gas rencana', '=COUNTBLANK(ringkasan_gas!E2:E366)'),
           ('Baris Realisasi produksi tanpa angka', '=COUNTIFS(harian_produksi!B:B,"Realisasi",harian_produksi!F:F,"")'),
           ('Unit tanpa RKAP tahunan', '=COUNTIF(rkap_bulanan!O:O,0)'),
           ('Safety stock / deadstock belum diisi', '=COUNTBLANK(parameter_kri!E2:F6)'),
           ('Rasio cons-rate belum diisi', '=COUNTBLANK(cons_rate!C2:C9)')]
    for j, h in enumerate(['Pemeriksaan', 'Hasil'], 2):
        gaya(ws.cell(r2 + 1, j, h), HEAD)
    for i, (a, f) in enumerate(CEK):
        gaya(ws.cell(r2 + 2 + i, 2, a), {'font': Font(name=F, size=10)})
        gaya(ws.cell(r2 + 2 + i, 3, f), RUMUS)
        ws.cell(r2 + 2 + i, 3).alignment = Alignment(horizontal='left')
    t(r2 + 3 + len(CEK), 2, 'Semua angka harus 0 dan posisi_data "OK" sebelum berkas dikirim.', size=9, italic=True, color='7F7F7F')
    ws.sheet_properties.tabColor = '1F5F3A'

    KELUAR.mkdir(exist_ok=True)
    out = KELUAR / f'TEMPLATE_Dalops_{kode}_2026.xlsx'
    wb.save(out)
    return out


if __name__ == '__main__':
    for k in ['PKT', 'PIM', 'PSP', 'PKC']:
        print('templat:', buat(k))
