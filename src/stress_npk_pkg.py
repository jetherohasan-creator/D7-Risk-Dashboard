#!/usr/bin/env python3
"""Stress test laju produksi NPK PT Petrokimia Gresik (PKG) untuk Goal Seek.

Pertanyaan: pada beban produksi NPK berapa stok bahan baku (ditambah pasokan
yang sudah terjadwal) tidak lagi mencukupi sampai akhir horizon?

Sumber angka: rencana harian Dalops PKG yang tertanam di dashboard v2beta
(objek D.sim, posisi 4 Oktober 2026, horizon 5 Okt – 31 Des 2026), yaitu
angka yang sama dengan modul 08 Stress Test di dashboard. Tiga bahan yang
mengikat pabrik NPK di dashboard: asam fosfat (di atas heel), KCl, amoniak.

Pakai:  python3 src/stress_npk_pkg.py  → analisis/StressTest_NPK_PKG.xlsx
"""
import datetime as dt
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AKAR = Path(__file__).resolve().parent.parent
ACUAN = AKAR / 'ref' / 'Dashboard_CAK_PRI_v2beta_04Oktober2026.html'
KELUAR = AKAR / 'analisis' / 'StressTest_NPK_PKG.xlsx'

F = 'Arial'
HEAD = dict(font=Font(name=F, bold=True, color='FFFFFF', size=10), fill=PatternFill('solid', fgColor='1F5F3A'),
            alignment=Alignment(horizontal='center', vertical='center', wrap_text=True))
SUB = dict(font=Font(name=F, bold=True, color='1F5F3A', size=10), fill=PatternFill('solid', fgColor='E2EFDA'),
           alignment=Alignment(horizontal='center', vertical='center', wrap_text=True))
ISI = dict(font=Font(name=F, color='0000FF', size=10), fill=PatternFill('solid', fgColor='FFF2CC'))
KUNCI = dict(font=Font(name=F, color='FF0000', size=10, bold=True), fill=PatternFill('solid', fgColor='FFFF00'))
DATA = dict(font=Font(name=F, color='000000', size=10))
RUMUS = dict(font=Font(name=F, color='000000', size=10))
LINK = dict(font=Font(name=F, color='008000', size=10))
HASIL = dict(font=Font(name=F, color='000000', size=11, bold=True), fill=PatternFill('solid', fgColor='DDEBF7'))
TEKS = dict(font=Font(name=F, size=10))
KECIL = dict(font=Font(name=F, size=9, color='595959'))
TIPIS = Side(style='thin', color='BFBFBF')
T, T1, PCT, TGL = '#,##0;(#,##0);-', '#,##0.0;(#,##0.0);-', '0.0%', 'dd-mmm-yyyy'


def gaya(c, g, fmt=None):
    for k, v in g.items():
        setattr(c, k, v)
    if fmt:
        c.number_format = fmt
    return c


def tulis(ws, ref, v, g=TEKS, fmt=None):
    return gaya(ws[ref] if isinstance(ref, str) else ws.cell(*ref), g, fmt) if v is None else gaya(_set(ws, ref, v), g, fmt)


def _set(ws, ref, v):
    c = ws[ref] if isinstance(ref, str) else ws.cell(*ref)
    c.value = v
    return c


def muat_D():
    teks = ACUAN.read_text(encoding='utf-8')
    m = re.search(r'^const D = (\{.*\});\s*$', teks, re.M)
    return json.loads(m.group(1))


def main():
    D = muat_D()
    S = D['sim']
    tgl = [dt.date.fromisoformat(t) for t in S['tgl']]
    N = len(tgl)
    unit = S['unit']
    wb = Workbook()

    # ============================ PANDUAN ============================
    ws = wb.active
    ws.title = 'PANDUAN'
    ws.sheet_view.showGridLines = False
    for c, w in zip('ABCDE', [3, 30, 70, 18, 18]):
        ws.column_dimensions[c].width = w
    tulis(ws, 'B1', 'Stress Test Laju Produksi NPK · PT Petrokimia Gresik', dict(font=Font(name=F, bold=True, size=16, color='1F5F3A')))
    tulis(ws, 'B2', f'Posisi data {D["tanggal"]} · horizon {tgl[0]:%d %b %Y} – {tgl[-1]:%d %b %Y} ({N} hari) · sumber: rencana harian Dalops PKG di dashboard CAK PRI v2beta', KECIL)
    tulis(ws, 'B4', 'Pertanyaan yang dijawab', dict(font=Font(name=F, bold=True, size=11)))
    tulis(ws, 'B5', 'Pada beban produksi NPK berapa stok bahan baku, ditambah pasokan yang sudah terjadwal, tidak lagi mencukupi kebutuhan pabrik NPK sampai akhir horizon?', TEKS)
    ws.merge_cells('B5:E5'); ws['B5'].alignment = Alignment(wrap_text=True); ws.row_dimensions[5].height = 30
    tulis(ws, 'B7', 'Cara memakai Goal Seek', dict(font=Font(name=F, bold=True, size=11)))
    langkah = [
        '1. Buka sheet Asumsi. Sel merah-kuning C5 = beban produksi NPK (100% = rencana Dalops). Inilah sel yang diubah Goal Seek.',
        '2. Atur skenario pasokan di Asumsi (sel kuning): produksi asam fosfat, kiriman PJA, impor, keterlambatan kargo KCl, produksi amoniak, dll.',
        '3. Data > What-If Analysis > Goal Seek:  Set cell = Hasil!C6  ·  To value = 0  ·  By changing cell = Asumsi!C5.',
        '4. Hasil Goal Seek di Asumsi!C5 adalah TITIK PATAH: beban tertinggi yang masih bisa dilayani stok sampai akhir horizon. Bahan yang mengikat tampil di Hasil!C8.',
        '5. Untuk titik peringatan (stok menyentuh safety stock): Set cell = Hasil!C7, To value = 0.',
        '6. Pembanding tanpa Goal Seek: Hasil!C12 menghitung titik patah secara langsung (rumus). Hasil Goal Seek harus sama dengan angka ini.',
        '7. Catat hasil tiap skenario di tabel Hasil (baris 26 dst.), lalu kembalikan Asumsi!C5 ke 100%.']
    for i, s in enumerate(langkah):
        c = tulis(ws, (8 + i, 2), s)
        ws.merge_cells(start_row=8 + i, start_column=2, end_row=8 + i, end_column=5)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[8 + i].height = 28
    tulis(ws, 'B16', 'Logika neraca harian', dict(font=Font(name=F, bold=True, size=11)))
    logika = [
        'Kebutuhan bahan untuk NPK pada hari t = pemakaian NPK rencana Dalops hari t × (produksi NPK simulasi ÷ produksi NPK rencana). Rasio bahan per ton NPK harian mengikuti bauran unit Phonska di Dalops.',
        'Stok siap pakai = stok kemarin + pasokan masuk − pemakaian di luar NPK − kebutuhan NPK, dikurangi batas bawah (heel/deadstock). Stok boleh negatif di model: angka negatif = kekurangan.',
        'Asam fosfat: masuk = produksi PA-I/II × (1 − rasio sludge) + kiriman PJA + impor. Produksi PA diasumsikan tidak dibatasi belerang/asam sulfat (untuk itu pakai modul Stress Test dashboard).',
        'KCl: masuk = kargo impor, bisa digeser sekian hari (keterlambatan). Amoniak: masuk = produksi Amoniak I A/I B + pengadaan luar; pemakaian lain = urea, ZA, penjualan.',
        'Titik patah = beban tertinggi sehingga stok siap pakai terendah dari ketiga bahan, sepanjang horizon, tepat 0.']
    for i, s in enumerate(logika):
        c = tulis(ws, (17 + i, 2), '• ' + s)
        ws.merge_cells(start_row=17 + i, start_column=2, end_row=17 + i, end_column=5)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[17 + i].height = 30
    tulis(ws, 'B23', 'Arti warna', dict(font=Font(name=F, bold=True, size=11)))
    for i, (g, t) in enumerate([(KUNCI, 'Sel yang diubah Goal Seek (beban NPK)'), (ISI, 'Asumsi skenario — boleh diubah'),
                                (DATA, 'Data rencana Dalops (teks hitam, sheet Rencana_Dalops)'),
                                (HASIL, 'Hasil utama'), (LINK, 'Teks hijau = mengambil nilai dari sheet lain')]):
        c = tulis(ws, (24 + i, 2), 'contoh', g)
        c.border = Border(TIPIS, TIPIS, TIPIS, TIPIS)
        tulis(ws, (24 + i, 3), t)
    tulis(ws, 'B30', 'Indikatif untuk penyaringan awal, bukan pengganti balans resmi Perencanaan dan Pengendalian Produksi.', dict(font=Font(name=F, italic=True, size=9, color='C00000')))

    # ============================ ASUMSI ============================
    a = wb.create_sheet('Asumsi')
    a.sheet_view.showGridLines = False
    for c, w in zip('ABCDE', [3, 46, 16, 14, 64]):
        a.column_dimensions[c].width = w
    tulis(a, 'B1', 'Asumsi & Tuas Skenario', dict(font=Font(name=F, bold=True, size=14, color='1F5F3A')))
    for j, h in enumerate(['Parameter', 'Nilai', 'Satuan', 'Sumber / keterangan'], 2):
        gaya(a.cell(3, j, h), HEAD)
    A = {}   # nama → alamat sel

    def baris(r, kunci, label, nilai, sat, ket, g=ISI, fmt=None):
        tulis(a, (r, 2), label)
        gaya(_set(a, (r, 3), nilai), g, fmt)
        tulis(a, (r, 4), sat, KECIL)
        c = tulis(a, (r, 5), ket, KECIL)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        A[kunci] = f'Asumsi!$C${r}'

    r = 4
    gaya(a.cell(r, 2, 'TUAS GOAL SEEK'), SUB); a.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    baris(5, 'beban', 'Beban produksi NPK (semua unit)', 1.0, '% rencana', 'Sel yang diubah Goal Seek. 100% = rencana harian Dalops; 110% = 10% di atas rencana.', KUNCI, PCT)
    baris(6, 'horizon', 'Akhir horizon uji', tgl[-1], 'tanggal', f'Maksimum {tgl[-1]:%d %b %Y} (akhir rencana Dalops). Dipendekkan untuk uji 30/60 hari.', ISI, TGL)
    baris(7, 'bebanTeknis', 'Beban teknis maksimum pabrik NPK', None, '% rencana', '[ISI] Batas atas kemampuan pabrik terhadap rencana. Hasil Goal Seek di atas angka ini berarti stok bukan pembatas.', ISI, PCT)
    r = 9
    gaya(a.cell(r, 2, 'BEBAN PER UNIT (opsional, dikalikan dengan beban di atas)'), SUB); a.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    for i, u in enumerate(unit):
        baris(10 + i, f'u{i}', u['nama'], 1.0, '% rencana',
              f'Rencana Dalops maks {max(u["renc"]):,.0f} t/hari; {sum(1 for v in u["renc"] if v > 1)} hari operasi dalam horizon. Isi 0% untuk unit yang dimatikan.'.replace(',', '.'), ISI, PCT)
    r = 10 + len(unit) + 1
    gaya(a.cell(r, 2, 'PASOKAN (100% = sesuai rencana Dalops)'), SUB); a.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    PAS = [('paProd', 'Produksi asam fosfat sendiri (PA-I, PA-II)', 1.0, '%', 'Dalops: produksi PA rencana. Tidak dibatasi belerang di model ini.'),
           ('paPja', 'Kiriman asam fosfat PJA', 1.0, '%', 'Petro Jordan Abadi.'),
           ('paImp', 'Impor asam fosfat', 1.0, '%', 'Kargo impor terjadwal.'),
           ('kclImp', 'Impor KCl', 1.0, '%', 'Volume kargo impor KCl terjadwal.'),
           ('kclTunda', 'Keterlambatan seluruh kargo KCl', 0, 'hari', 'Menggeser kedatangan KCl sekian hari; kargo yang lewat akhir horizon hilang dari hitungan.'),
           ('nh3Prod', 'Produksi amoniak sendiri (I A, I B)', 1.0, '%', 'Bergantung pasokan gas.'),
           ('nh3Luar', 'Pengadaan amoniak dari luar', 1.0, '%', 'Anper PI Group dan sumber luar.'),
           ('nh3Lain', 'Pemakaian amoniak di luar NPK', 1.0, '%', 'Urea, ZA, penjualan. Turunkan untuk menguji prioritas amoniak ke NPK.'),
           ('paLain', 'Pemakaian asam fosfat di luar NPK', 1.0, '%', 'SP-36 dan lainnya.')]
    for i, (k, lbl, v, sat, ket) in enumerate(PAS):
        baris(r + 1 + i, k, lbl, v, sat, ket, ISI, PCT if sat == '%' else '0')
    r = r + 1 + len(PAS) + 1
    gaya(a.cell(r, 2, 'STOK AWAL & BATAS (posisi data)'), SUB); a.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    s0 = S['s0']
    STOK = [('s0PA', 'Stok awal asam fosfat (cair, basis P2O5 54%)', s0['PAliq'], 't', f'Dalops, posisi {D["tanggal"]}. Sudah bersih dari sludge.', DATA, T),
            ('s0KCL', 'Stok awal KCl (merah + putih)', s0['KCL'], 't', f'Dalops, posisi {D["tanggal"]}.', DATA, T),
            ('s0NH3', 'Stok awal amoniak', s0['NH3'], 't', f'Dalops, posisi {D["tanggal"]}.', DATA, T),
            ('lbPA', 'Batas bawah asam fosfat (heel)', S['heel'], 't', 'Heel tangki yang tidak dapat dipompa, sama dengan simulator dashboard. Ganti ke 6.000 bila ingin memakai deadstock Kesepakatan.', ISI, T),
            ('lbKCL', 'Batas bawah KCl (deadstock)', S['dead']['KCL'], 't', 'Deadstock Kesepakatan PI Group. Isi 0 untuk menyamakan dengan simulator dashboard.', ISI, T),
            ('lbNH3', 'Batas bawah amoniak (deadstock)', S['dead']['NH3'], 't', 'Deadstock Kesepakatan PI Group. Isi 0 untuk menyamakan dengan simulator dashboard.', ISI, T),
            ('ssPA', 'Safety stock asam fosfat', S['ss']['PA'], 't', 'Kesepakatan Safety Stock Bahan Baku PI Group.', DATA, T),
            ('ssKCL', 'Safety stock KCl', S['ss']['KCL'], 't', 'Kesepakatan Safety Stock Bahan Baku PI Group.', DATA, T),
            ('ssNH3', 'Safety stock amoniak', S['ss']['NH3'], 't', 'Angka internal PKG (tidak tercantum di Kesepakatan PI Group).', DATA, T)]
    for i, (k, lbl, v, sat, ket, g, fmt) in enumerate(STOK):
        baris(r + 1 + i, k, lbl, v, sat, ket, g, fmt)
    a.freeze_panes = 'A4'

    # ============================ RENCANA_DALOPS ============================
    d = wb.create_sheet('Rencana_Dalops')
    kol = (['Tanggal'] + [u['nama'] for u in unit] + ['NPK rencana (t)',
           'PA produksi', 'Rasio sludge', 'PA kiriman PJA', 'PA impor', 'PA pakai lain', 'PA untuk NPK',
           'KCl impor', 'KCl pakai lain', 'KCl untuk NPK',
           'NH3 produksi', 'NH3 luar', 'NH3 pakai lain', 'NH3 untuk NPK'])
    grup = [('', 1), ('PRODUKSI NPK RENCANA (t/hari)', len(unit) + 1), ('ASAM FOSFAT (t/hari)', 6), ('KCl (t/hari)', 3), ('AMONIAK (t/hari)', 4)]
    c0 = 1
    for g, n in grup:
        if g:
            gaya(d.cell(1, c0, g), SUB); d.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + n - 1)
        c0 += n
    for j, h in enumerate(kol, 1):
        gaya(d.cell(2, j, h), HEAD)
        d.column_dimensions[get_column_letter(j)].width = 13 if j > 1 else 13
    d.row_dimensions[2].height = 32
    seri = ([u['renc'] for u in unit] + [None, S['pa']['prod'], S['r']['sl'], S['pa']['pja'], S['pa']['imp'], S['pa']['lain'],
            S['pa']['uNPK'], S['kcl']['imp'], S['kcl']['lain'], S['kcl']['uNPK'], S['nh3']['prod'], S['nh3']['luar'],
            S['nh3']['lain'], S['nh3']['uNPK']])
    nU = len(unit)
    for i in range(N):
        r = 3 + i
        gaya(d.cell(r, 1, tgl[i]), DATA, TGL)
        for j, s in enumerate(seri, 2):
            if s is None:
                gaya(d.cell(r, j, f'=SUM(B{r}:{get_column_letter(1 + nU)}{r})'), RUMUS, T1)
            else:
                gaya(d.cell(r, j, round(s[i], 10)), DATA, '0.0000' if kol[j - 1] == 'Rasio sludge' else T1)
    rT = 3 + N
    gaya(d.cell(rT, 1, 'Jumlah'), dict(font=Font(name=F, bold=True, size=10)))
    for j in range(2, len(kol) + 1):
        if kol[j - 1] != 'Rasio sludge':
            gaya(d.cell(rT, j, f'=SUM({get_column_letter(j)}3:{get_column_letter(j)}{rT - 1})'), dict(font=Font(name=F, bold=True, size=10)), T)
    tulis(d, (rT + 2, 1), f'Sumber: rencana harian Dalops PKG (D.sim) yang tertanam di Dashboard CAK PRI v2beta, posisi data {D["tanggal"]}; '
          'sama dengan angka simulator modul 08 di dashboard. Nilai hardcode, bukan rumus; perbarui dari Dalops terbaru bila posisi data berganti.', KECIL)
    d.freeze_panes = 'B3'
    L = {k: get_column_letter(i + 1) for i, k in enumerate(kol)}       # kolom per judul

    # ============================ NERACA ============================
    n = wb.create_sheet('Neraca')
    hdr = ['Tanggal', 'Hari ke', 'Dalam horizon', 'NPK simulasi (t)', 'NPK rencana (t)',
           'PA masuk', 'PA keluar', 'PA stok siap pakai', 'KCl masuk', 'KCl keluar', 'KCl stok siap pakai',
           'NH3 masuk', 'NH3 keluar', 'NH3 stok siap pakai',
           'PA kebutuhan @100% kum.', 'PA ruang kum.', 'KCl kebutuhan @100% kum.', 'KCl ruang kum.', 'NH3 kebutuhan @100% kum.', 'NH3 ruang kum.',
           'Beban maks PA s.d. hari ini', 'Beban maks KCl s.d. hari ini', 'Beban maks NH3 s.d. hari ini',
           'PA habis (tgl)', 'KCl habis (tgl)', 'NH3 habis (tgl)', 'Stok − SS terendah hari ini',
           'PA siap (horizon)', 'KCl siap (horizon)', 'NH3 siap (horizon)']
    grupN = [('', 5), ('ASAM FOSFAT (t)', 3), ('KCl (t)', 3), ('AMONIAK (t)', 3), ('PEMBANDING ANALITIK (untuk cek Goal Seek)', 9), ('PENANDA', 7)]
    c0 = 1
    for g, k in grupN:
        if g:
            gaya(n.cell(1, c0, g), SUB); n.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + k - 1)
        c0 += k
    for j, h in enumerate(hdr, 1):
        gaya(n.cell(2, j, h), HEAD)
        n.column_dimensions[get_column_letter(j)].width = 12.5
    n.row_dimensions[2].height = 44
    komen = {8: 'Stok di atas heel. Negatif = kekurangan (model tidak memotong produksi, supaya Goal Seek bisa mencari titik patah).',
             16: 'Stok awal − batas bawah + kumulatif pasokan − kumulatif pemakaian lain. Ruang yang tersedia untuk NPK sampai hari ini.',
             21: 'Beban tertinggi yang masih aman sampai hari ini = ruang kumulatif ÷ kebutuhan kumulatif pada beban 100%.'}
    for j, t in komen.items():
        n.cell(2, j).comment = Comment(t, 'CAK PRI')
    sumU = '+'.join(f"Rencana_Dalops!{get_column_letter(2 + i)}{{r}}*{A[f'u{i}']}" for i in range(nU))
    for i in range(N):
        r = 3 + i
        rd = 3 + i                            # baris yang sama di Rencana_Dalops
        prev = r - 1
        rasio = f'IF(E{r}>0,D{r}/E{r},0)'
        f = {}
        f['A'] = (f'=Rencana_Dalops!A{rd}', LINK, TGL)
        f['B'] = (f'=A{r}-Rencana_Dalops!$A$3+1', RUMUS, '0')
        f['C'] = (f'=IF(A{r}<={A["horizon"]},1,0)', RUMUS, '0')
        f['D'] = (f'=({sumU.format(r=rd)})*{A["beban"]}', RUMUS, T1)
        f['E'] = (f'=Rencana_Dalops!{L["NPK rencana (t)"]}{rd}', LINK, T1)
        # asam fosfat
        f['F'] = (f'=Rencana_Dalops!{L["PA produksi"]}{rd}*{A["paProd"]}*(1-Rencana_Dalops!{L["Rasio sludge"]}{rd})'
                  f'+Rencana_Dalops!{L["PA kiriman PJA"]}{rd}*{A["paPja"]}+Rencana_Dalops!{L["PA impor"]}{rd}*{A["paImp"]}', RUMUS, T1)
        f['G'] = (f'=Rencana_Dalops!{L["PA pakai lain"]}{rd}*{A["paLain"]}+Rencana_Dalops!{L["PA untuk NPK"]}{rd}*{rasio}', RUMUS, T1)
        f['H'] = ((f'={A["s0PA"]}-{A["lbPA"]}+F{r}-G{r}' if i == 0 else f'=H{prev}+F{r}-G{r}'), RUMUS, T)
        # KCl (kargo bisa digeser)
        kclIn = (f'IF(B{r}-{A["kclTunda"]}>=1,INDEX(Rencana_Dalops!${L["KCl impor"]}$3:${L["KCl impor"]}${2 + N},B{r}-{A["kclTunda"]}),0)')
        f['I'] = (f'={kclIn}*{A["kclImp"]}', RUMUS, T1)
        f['J'] = (f'=Rencana_Dalops!{L["KCl pakai lain"]}{rd}+Rencana_Dalops!{L["KCl untuk NPK"]}{rd}*{rasio}', RUMUS, T1)
        f['K'] = ((f'={A["s0KCL"]}-{A["lbKCL"]}+I{r}-J{r}' if i == 0 else f'=K{prev}+I{r}-J{r}'), RUMUS, T)
        # amoniak
        f['L'] = (f'=Rencana_Dalops!{L["NH3 produksi"]}{rd}*{A["nh3Prod"]}+Rencana_Dalops!{L["NH3 luar"]}{rd}*{A["nh3Luar"]}', RUMUS, T1)
        f['M'] = (f'=Rencana_Dalops!{L["NH3 pakai lain"]}{rd}*{A["nh3Lain"]}+Rencana_Dalops!{L["NH3 untuk NPK"]}{rd}*{rasio}', RUMUS, T1)
        f['N'] = ((f'={A["s0NH3"]}-{A["lbNH3"]}+L{r}-M{r}' if i == 0 else f'=N{prev}+L{r}-M{r}'), RUMUS, T)
        # pembanding analitik: kebutuhan kumulatif pada beban global 100% dan ruang kumulatif
        npk100 = f'IF({A["beban"]}>0,{rasio}/{A["beban"]},0)'
        f['O'] = ((f'=' if i == 0 else f'=O{prev}+') + f'Rencana_Dalops!{L["PA untuk NPK"]}{rd}*{npk100}', RUMUS, T)
        f['P'] = (f'=H{r}+O{r}*{A["beban"]}', RUMUS, T)
        f['Q'] = ((f'=' if i == 0 else f'=Q{prev}+') + f'Rencana_Dalops!{L["KCl untuk NPK"]}{rd}*{npk100}', RUMUS, T)
        f['R'] = (f'=K{r}+Q{r}*{A["beban"]}', RUMUS, T)
        f['S'] = ((f'=' if i == 0 else f'=S{prev}+') + f'Rencana_Dalops!{L["NH3 untuk NPK"]}{rd}*{npk100}', RUMUS, T)
        f['T'] = (f'=N{r}+S{r}*{A["beban"]}', RUMUS, T)
        f['U'] = (f'=IF(OR(C{r}=0,O{r}<=0),"",P{r}/O{r})', RUMUS, PCT)
        f['V'] = (f'=IF(OR(C{r}=0,Q{r}<=0),"",R{r}/Q{r})', RUMUS, PCT)
        f['W'] = (f'=IF(OR(C{r}=0,S{r}<=0),"",T{r}/S{r})', RUMUS, PCT)
        f['X'] = (f'=IF(AND(C{r}=1,H{r}<0),A{r},"")', RUMUS, TGL)
        f['Y'] = (f'=IF(AND(C{r}=1,K{r}<0),A{r},"")', RUMUS, TGL)
        f['Z'] = (f'=IF(AND(C{r}=1,N{r}<0),A{r},"")', RUMUS, TGL)
        f['AA'] = (f'=IF(C{r}=0,"",MIN(H{r}+{A["lbPA"]}-{A["ssPA"]},K{r}+{A["lbKCL"]}-{A["ssKCL"]},N{r}+{A["lbNH3"]}-{A["ssNH3"]}))', RUMUS, T)
        f['AB'] = (f'=IF(C{r}=1,H{r},"")', RUMUS, T)
        f['AC'] = (f'=IF(C{r}=1,K{r},"")', RUMUS, T)
        f['AD'] = (f'=IF(C{r}=1,N{r},"")', RUMUS, T)
        for col, (v, g, fmt) in f.items():
            gaya(n[f'{col}{r}'], g, fmt).value = v
    rN = 2 + N
    n.freeze_panes = 'D3'
    rng = lambda c: f'Neraca!${c}$3:${c}${rN}'

    # ============================ HASIL ============================
    h = wb.create_sheet('Hasil', 1)
    h.sheet_view.showGridLines = False
    for c, w in zip('ABCDEFG', [3, 74, 18, 18, 18, 18, 30]):
        h.column_dimensions[c].width = w
    tulis(h, 'B1', 'Hasil Stress Test Laju Produksi NPK · PKG', dict(font=Font(name=F, bold=True, size=14, color='1F5F3A')))
    tulis(h, 'B2', '="Beban diuji: "&ROUND(Asumsi!C5*100,1)&"% · horizon s.d. "&DAY(Asumsi!C6)&"/"&MONTH(Asumsi!C6)&"/"&YEAR(Asumsi!C6)&" · indikatif, bukan pengganti balans resmi"', KECIL)
    gaya(h.cell(4, 2, 'SEL TARGET GOAL SEEK'), SUB); h.merge_cells('B4:D4')
    tulis(h, 'B5', 'Ukuran')
    gaya(h['C5'], HEAD).value = 'Nilai'
    tulis(h, 'B6', 'Stok siap pakai terendah dari 3 bahan selama horizon (t)  →  Goal Seek ke 0 = TITIK PATAH', dict(font=Font(name=F, bold=True, size=10)))
    gaya(h['C6'], HASIL, T).value = '=MIN(D16:F16)'
    tulis(h, 'B7', 'Selisih stok terhadap safety stock, terendah (t)  →  Goal Seek ke 0 = TITIK PERINGATAN', dict(font=Font(name=F, bold=True, size=10)))
    gaya(h['C7'], HASIL, T).value = f'=MIN({rng("AA")})'
    tulis(h, 'B8', 'Bahan baku pengikat (stok siap pakai paling tipis)')
    gaya(h['C8'], HASIL).value = '=INDEX(D14:F14,MATCH(MIN(D16:F16),D16:F16,0))'
    tulis(h, 'B9', 'Status pada beban diuji')
    gaya(h['C9'], HASIL).value = '=IF(C6<-0.5,"STOK TIDAK CUKUP",IF(C6<=0.5,"TEPAT DI TITIK PATAH","Stok cukup"))'

    gaya(h.cell(10, 2, 'TITIK PATAH TANPA GOAL SEEK (rumus langsung, untuk mencocokkan)'), SUB); h.merge_cells('B10:D10')
    tulis(h, 'B11', 'Beban maksimum yang masih dicukupi stok, per bahan & keseluruhan')
    tulis(h, 'B12', 'Titik patah = beban maksimum terkecil', dict(font=Font(name=F, bold=True, size=10)))
    gaya(h['C12'], HASIL, PCT).value = '=MIN(D17:F17)'
    tulis(h, 'B13', 'Setara produksi NPK rata-rata pada titik patah (t/hari)')
    gaya(h['C13'], HASIL, T).value = f'=C12*SUMPRODUCT({rng("E")},{rng("C")})/SUM({rng("C")})'

    for j, nm in enumerate(['Asam fosfat', 'KCl', 'Amoniak'], 4):
        gaya(h.cell(14, j, nm), HEAD)
    tulis(h, 'B14', 'Per bahan baku', dict(font=Font(name=F, bold=True, size=10)))
    BAR = [('Stok siap pakai awal (di atas batas bawah, t)', [f'={A["s0PA"]}-{A["lbPA"]}', f'={A["s0KCL"]}-{A["lbKCL"]}', f'={A["s0NH3"]}-{A["lbNH3"]}'], T),
           ('Stok siap pakai terendah pada beban diuji (t)', [f'=MIN({rng(c)})' for c in ('AB', 'AC', 'AD')], T),
           ('Beban maksimum yang masih dicukupi', [f'=MIN({rng(c)})' for c in 'UVW'], PCT),
           ('Tanggal pertama stok habis pada beban diuji', [f'=IF(MIN({rng(c)})=0,"tidak habis",MIN({rng(c)}))' for c in 'XYZ'], TGL),
           ('Stok siap pakai di akhir horizon (t)', [f'=INDEX({rng(c)},SUM({rng("C")}))' for c in 'HKN'], T),
           ('Total pasokan masuk dalam horizon (t)', [f'=SUMPRODUCT({rng(c)},{rng("C")})' for c in 'FIL'], T),
           ('Total kebutuhan NPK + lain dalam horizon (t)', [f'=SUMPRODUCT({rng(c)},{rng("C")})' for c in 'GJM'], T)]
    for i, (lbl, fs, fmt) in enumerate(BAR):
        r = 15 + i
        tulis(h, (r, 2), lbl)
        for j, fx in enumerate(fs, 4):
            gaya(h.cell(r, j, fx), RUMUS if i else LINK, fmt)
    h['C15'].value = None
    tulis(h, 'B24', 'Artinya', dict(font=Font(name=F, bold=True, size=10)))
    c = gaya(h['C24'], HASIL)
    c.value = ('=IF(C12<0,"Stok "&INDEX(D14:F14,MATCH(C12,D17:F17,0))&" habis bahkan bila pabrik NPK berhenti: pemakaian di luar NPK (urea, ZA, penjualan) harus dikurangi lebih dulu.",'
               'IF(C12<1,"Rencana NPK tidak bisa dipenuhi penuh. Beban maksimum "&ROUND(C12*100,1)&"%"&" dari rencana; pengikat: "&INDEX(D14:F14,MATCH(C12,D17:F17,0))&".",'
               '"Rencana NPK tercukupi. Masih ada ruang "&ROUND((C12-1)*100,1)&"%"&" di atas rencana sebelum "&INDEX(D14:F14,MATCH(C12,D17:F17,0))&" menjadi pembatas."))')
    h.merge_cells('C24:G24'); c.alignment = Alignment(wrap_text=True, vertical='top'); h.row_dimensions[24].height = 32
    tulis(h, 'B23', '="Bila titik patah di atas beban teknis maksimum ("&IF(Asumsi!C7="","[ISI]",ROUND(Asumsi!C7*100,0)&"%")&"), stok bahan baku bukan pembatas dalam horizon ini."', KECIL)

    gaya(h.cell(25, 2, 'CATATAN HASIL PER SKENARIO (isi manual setelah Goal Seek)'), SUB); h.merge_cells('B25:G25')
    for j, t in enumerate(['Skenario', 'Titik patah (beban)', 'Bahan pengikat', 'Tgl habis @100%', 'Titik peringatan', 'Catatan'], 2):
        gaya(h.cell(26, j, t), HEAD)
    contoh = ['Rencana dasar (semua 100%)', 'KCl mundur 30 hari', 'PJA turun 50%', 'Produksi amoniak −30%', 'Impor PA batal']
    for i, s in enumerate(contoh):
        tulis(h, (27 + i, 2), s)
        for j in range(3, 8):
            gaya(h.cell(27 + i, j), ISI, PCT if j in (3, 6) else None)

    # grafik stok siap pakai
    ch = LineChart()
    ch.title = 'Stok siap pakai pada beban diuji (t)'
    ch.y_axis.title = 'ton'
    ch.height, ch.width = 8, 22
    for c in (8, 11, 14):
        ch.add_data(Reference(n, min_col=c, min_row=2, max_row=rN), titles_from_data=True)
    ch.set_categories(Reference(n, min_col=1, min_row=3, max_row=rN))
    ch.x_axis.number_format = 'dd-mmm'
    h.add_chart(ch, 'B34')

    wb._sheets = [wb['PANDUAN'], wb['Hasil'], wb['Asumsi'], wb['Neraca'], wb['Rencana_Dalops']]
    KELUAR.parent.mkdir(exist_ok=True)
    wb.save(KELUAR)
    print('stress test:', KELUAR, f'({N} hari, {nU} unit)')


if __name__ == '__main__':
    main()
