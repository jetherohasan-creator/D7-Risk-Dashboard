"""Uji pulang-pergi templat: isi templat dengan angka DUMMY, lalu bandingkan di dashboard.
Pakai: python3 tests/isi_templat.py PKT  → tests/_tangkapan/TERISI_<KODE>.xlsx"""
import sys, json, datetime as dt
from pathlib import Path
import openpyxl
A = Path(__file__).resolve().parent.parent
kode = sys.argv[1]
J = json.loads((A/'data/DUMMY_CAK_PRI_PIGroup_data.json').read_text())['entitas'][kode]
R = lambda s: [dict(zip(J[s]['kolom'], r)) for r in J[s]['baris']]
D = lambda s: dt.datetime.strptime(s, '%Y-%m-%d') if isinstance(s, str) else s
wb = openpyxl.load_workbook(A/f'templates/TEMPLATE_Dalops_{kode}_2026.xlsx')
pr = {r['parameter']: r['nilai'] for r in R('profil')}
ws = wb['profil']
for r in range(2, ws.max_row+1):
    k = ws.cell(r, 1).value
    if k in ('posisi_data', 'tanggal_update'): ws.cell(r, 2).value = D(pr[k])
    elif k in ('konteks', 'dosis_urea_kg_per_ha', 'dosis_npk_kg_per_ha', 'produk_npk', 'warna_aksen'): ws.cell(r, 2).value = pr[k]
def isi(sheet, kunci, kolom, sumber, override=()):
    ws = wb[sheet]; hdr = [c.value for c in ws[1]]; idx = {h: i+1 for i, h in enumerate(hdr)}
    peta = {tuple((str(r[k])[:10] if k == 'tanggal' else str(r[k])) for k in kunci): r for r in sumber}
    for r in range(2, ws.max_row+1):
        key = tuple((ws.cell(r, idx[k]).value.strftime('%Y-%m-%d') if isinstance(ws.cell(r, idx[k]).value, dt.datetime) else str(ws.cell(r, idx[k]).value)) for k in kunci)
        s = peta.get(key)
        if not s: continue
        for k in kolom:
            c = ws.cell(r, idx[k])
            if k in override or not (isinstance(c.value, str) and c.value.startswith('=')): c.value = s[k]
isi('harian_gas', ['tanggal', 'pemasok'], ['pasokan_bbtud'], R('harian_gas'))
isi('ringkasan_gas', ['tanggal'], ['kebutuhan_rencana_bbtud', 'pemakaian_aktual_bbtud', 'gas_per_ton_amoniak_mmbtu', 'gas_per_ton_urea_mmbtu'], R('ringkasan_gas'))
isi('harian_bahan_baku', ['tanggal', 'kode_material'], ['stok_awal_t', 'kedatangan_t', 'produksi_sendiri_t', 'pemakaian_t', 'dijual_transfer_t', 'stok_akhir_t'], R('harian_bahan_baku'), override=('stok_akhir_t',))
isi('harian_produksi', ['tanggal', 'unit'], ['kapasitas_t_per_hari', 'produksi_t'], R('harian_produksi'))
isi('rkap_bulanan', ['unit'], ['jan','feb','mar','apr','mei','jun','jul','agu','sep','okt','nov','des'], R('rkap_bulanan'))
isi('stok_produk_lini1', ['tanggal', 'produk'], ['stok_awal_t', 'masuk_t', 'penyaluran_t', 'stok_akhir_t', 'curah_di_pabrik_t'], R('stok_produk_lini1'), override=('stok_akhir_t',))
isi('status_unit', ['unit'], ['status', 'mulai', 'selesai', 'keterangan'], [dict(r, mulai=D(r['mulai']), selesai=D(r['selesai'])) for r in R('status_unit')])
isi('parameter_kri', ['kode_material'], ['safety_stock', 'deadstock'], R('parameter_kri'))
isi('cons_rate', ['bahan'], ['rasio'], R('cons_rate'))
isi('sumber_pasokan', ['id'], ['volume_acuan'], R('sumber_pasokan'))
def tabel(sheet, sumber, tgl):
    ws = wb[sheet]; hdr = [c.value for c in ws[1]]
    for i, s in enumerate(sumber, 2):
        for j, h in enumerate(hdr, 1):
            if h in s and not (isinstance(ws.cell(i, j).value, str) and str(ws.cell(i, j).value).startswith('=')):
                ws.cell(i, j).value = D(s[h]) if h in tgl else s[h]
tabel('jadwal_kapal', R('jadwal_kapal'), ('eta_terkini', 'eta_awal'))
tabel('perubahan_versi', R('perubahan_versi'), ('versi_lama', 'versi_baru', 'eta_lama', 'eta_baru'))
tabel('transfer_antar_entitas', R('transfer_antar_entitas'), ('tanggal',))
out = A/f'tests/_tangkapan/TERISI_{kode}.xlsx'; out.parent.mkdir(exist_ok=True); wb.save(out); print(out)
