"""Uji Excel stress test NPK PKG (analisis/StressTest_NPK_PKG.xlsx).
1. Beban 100% dengan batas bawah = simulator dashboard → neraca sama dengan simulator v2beta.
2. Beban diisi titik patah analitik → stok siap pakai terendah = 0 (setara hasil Goal Seek).
3. Beberapa skenario: titik patah dan bahan pengikatnya.
Pakai: python3 tests/stress_npk.py <path recalc.py>"""
import json, re, subprocess, sys, tempfile, shutil
from pathlib import Path
import openpyxl

A = Path(__file__).resolve().parent.parent
XL = A / 'analisis' / 'StressTest_NPK_PKG.xlsx'
RECALC = sys.argv[1]
TMP = Path(tempfile.mkdtemp())
D = json.loads(re.search(r'^const D = (\{.*\});\s*$', (A / 'ref/Dashboard_CAK_PRI_v2beta_04Oktober2026.html').read_text(), re.M).group(1))
Q = D['sim']; n = len(Q['tgl'])


def jalankan(nama, ubah):
    wb = openpyxl.load_workbook(XL); a = wb['Asumsi']
    lab = {a.cell(r, 2).value: r for r in range(1, a.max_row + 1)}
    for k, v in ubah.items():
        a.cell(lab[k], 3).value = v
    f = TMP / f'{nama}.xlsx'; wb.save(f)
    out = json.loads(subprocess.run([sys.executable, RECALC, str(f), '120'], capture_output=True, text=True).stdout)
    assert out.get('total_errors') == 0, out
    return openpyxl.load_workbook(f, data_only=True)


def sim_dashboard():
    SUL, SA, PA, PR, KCL, NH3 = (Q['s0'][k] for k in ('SUL', 'SA', 'PAliq', 'PR', 'KCL', 'NH3'))
    out = []
    for i in range(n):
        SUL = max(SUL + Q['sul']['imp'][i] + Q['sul']['lain'][i], 0); rs = Q['r']['sa'][i]
        sa2 = min(Q['sa']['pSA2'][i], SUL / rs if rs > 0 else 0); SUL = max(0, SUL - sa2 * rs)
        SA += Q['sa']['pSA1'][i] + sa2 + sum(Q['sa'][k][i] for k in ('smelt', 'amman', 'manyar', 'imp'))
        PR += Q['pr']['imp'][i]; KCL += Q['kcl']['imp'][i]; NH3 += Q['nh3']['prod'][i] + Q['nh3']['luar'][i]
        KCL -= min(Q['kcl']['lain'][i], KCL); NH3 -= min(Q['nh3']['lain'][i], NH3)
        rsa, rpr, rsl = Q['r']['pasa'][i], Q['r']['papr'][i], Q['r']['sl'][i]; renc = Q['pa']['prod'][i]; lainSA = Q['sa']['lain'][i]
        need = lainSA + renc * rsa; fSA = min(1, SA / need) if need > 0 else 1
        pa = min(renc, renc * fSA, PR / rpr if rpr > 0 else 0)
        SA = max(0, SA - lainSA * fSA - pa * rsa); PR = max(0, PR - pa * rpr)
        PA += pa + Q['pa']['pja'][i] + Q['pa']['imp'][i] - pa * rsl
        PA -= min(Q['pa']['lain'][i], max(0, PA - Q['heel']))
        rP, rK, rN = Q['pa']['uNPK'][i], Q['kcl']['uNPK'][i], Q['nh3']['uNPK'][i]
        f = min(1, max(0, PA - Q['heel']) / rP if rP else 1, KCL / rK if rK else 1, NH3 / rN if rN else 1)
        PA -= rP * f; KCL -= rK * f; NH3 -= rN * f
        out.append((PA - Q['heel'], KCL, NH3))
    return out


gagal = 0
wb = jalankan('dasar0', {'Batas bawah KCl (deadstock)': 0, 'Batas bawah amoniak (deadstock)': 0})
ws = wb['Neraca']; ds = sim_dashboard()
mx = max(abs(ws.cell(3 + i, c).value - ds[i][k]) for i in range(n) for k, c in enumerate((8, 11, 14)))
print(f'1. Beban 100% vs simulator dashboard: selisih maks {mx:.4f} t'); gagal += mx > 0.01

SKEN = [('Rencana dasar', {}), ('KCl mundur 30 hari', {'Keterlambatan seluruh kargo KCl': 30}),
        ('PJA turun 50%', {'Kiriman asam fosfat PJA': 0.5}), ('Produksi amoniak −30%', {'Produksi amoniak sendiri (I A, I B)': 0.7}),
        ('Impor asam fosfat batal', {'Impor asam fosfat': 0}),
        ('Gabungan: PJA −50% + amoniak −30%', {'Kiriman asam fosfat PJA': 0.5, 'Produksi amoniak sendiri (I A, I B)': 0.7})]
print('\n2–3. Titik patah per skenario (batas bawah = heel PA, deadstock KCl & amoniak)')
for i, (nm, u) in enumerate(SKEN):
    h = jalankan(f's{i}', u)['Hasil']
    tp = h['C12'].value
    h2 = jalankan(f's{i}b', dict(u, **{'Beban produksi NPK (semua unit)': tp}))['Hasil']
    ok = abs(h2['C6'].value) < 0.5
    gagal += not ok
    tgl = [h[f'{c}18'].value for c in 'DEF']
    print(f'   {nm:38s} titik patah {tp*100:6.1f}% · pengikat {h2["C8"].value:11s} · '
          f'@100%: {h["C9"].value}, habis {", ".join(str(x)[:10] for x in tgl)} · cek Goal Seek C6={h2["C6"].value:+.2f} t {"OK" if ok else "GAGAL"}')
shutil.rmtree(TMP)
sys.exit(1 if gagal else 0)
